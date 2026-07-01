from __future__ import annotations

import logging

from fastapi import APIRouter, File, HTTPException, Response, UploadFile, status

from briefcase import repository as repo
from briefcase.ingest.parsers import supported_extensions
from briefcase.ingest.service import IngestionError, IngestionService
from briefcase.models import (
    AddTextSource,
    AddUrlSource,
    Answer,
    ChatRequest,
    CreateNotebook,
    Draft,
    DraftCommand,
    DraftEdit,
    DraftVersion,
    Notebook,
    SaveDraft,
    Source,
)
from briefcase.rag.author import DraftAuthor
from briefcase.rag.generator import AnswerGenerator
from briefcase.rag.llm import OllamaClient, llm_status
from briefcase.rag import studio
from briefcase import settings_store
from briefcase.ingest import ocr

router = APIRouter(prefix="/api")

_ingestion = IngestionService()
_generator = AnswerGenerator()
_author = DraftAuthor()


# ---------------- system ----------------


@router.get("/status")
def status_info() -> dict:
    from briefcase.config import get_settings

    cfg = get_settings()
    provider = settings_store.get("embeddings_provider") or cfg.embeddings.provider
    return {
        "llm": llm_status(),
        "embeddings": provider,
        "ocr": ocr.status(),
        "supported_extensions": sorted(supported_extensions()),
    }


# ---------------- settings ----------------


@router.get("/settings")
def get_settings_view() -> dict:
    """Current settings for the UI, plus what is actually active/available."""
    ollama_url = settings_store.get("ollama_url") or get_settings_default("ollama_url")
    return {
        "settings": settings_store.public_view(),
        "active": {"llm": llm_status(), "ocr": ocr.status()},
        "available": {
            "ollama": OllamaClient.is_available(ollama_url),
            "ollama_url": ollama_url,
        },
        "options": {
            "llm_provider": ["auto", "gemini", "ollama", "extractive"],
            "embeddings_provider": ["hashing", "sentence-transformers", "gemini"],
            "ocr": ["auto", "off"],
        },
    }


@router.put("/settings")
def update_settings(patch: dict) -> dict:
    settings_store.update(patch)
    return get_settings_view()


@router.delete("/settings/{field}")
def delete_setting(field: str) -> dict:
    settings_store.delete(field)
    return get_settings_view()


@router.post("/settings/test-ollama")
def test_ollama(payload: dict | None = None) -> dict:
    url = (payload or {}).get("url") or settings_store.get("ollama_url") \
        or get_settings_default("ollama_url")
    ok = OllamaClient.is_available(url)
    models: list[str] = []
    if ok:
        import json as _json
        import urllib.request

        try:
            with urllib.request.urlopen(f"{url.rstrip('/')}/api/tags", timeout=2) as r:
                data = _json.loads(r.read().decode("utf-8"))
                models = [m.get("name", "") for m in data.get("models", [])]
        except Exception:
            pass
    return {"reachable": ok, "url": url, "models": models}


def get_settings_default(name: str) -> str:
    from briefcase.config import get_settings

    return getattr(get_settings().llm, name)


# ---------------- notebooks ----------------


@router.get("/notebooks", response_model=list[Notebook])
def list_notebooks() -> list[Notebook]:
    return repo.list_notebooks()


@router.post("/notebooks", response_model=Notebook, status_code=status.HTTP_201_CREATED)
def create_notebook(payload: CreateNotebook) -> Notebook:
    return repo.create_notebook(payload.title, payload.description)


@router.get("/notebooks/{notebook_id}", response_model=Notebook)
def get_notebook(notebook_id: str) -> Notebook:
    nb = repo.get_notebook(notebook_id)
    if nb is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Notebook not found")
    return nb


@router.delete("/notebooks/{notebook_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_notebook(notebook_id: str) -> None:
    if not repo.delete_notebook(notebook_id):
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Notebook not found")


# ---------------- sources ----------------


def _require_notebook(notebook_id: str) -> None:
    if repo.get_notebook(notebook_id) is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Notebook not found")


@router.get("/notebooks/{notebook_id}/sources", response_model=list[Source])
def list_sources(notebook_id: str) -> list[Source]:
    _require_notebook(notebook_id)
    return repo.list_sources(notebook_id)


@router.post(
    "/notebooks/{notebook_id}/sources/upload",
    response_model=Source,
    status_code=status.HTTP_201_CREATED,
)
async def upload_source(notebook_id: str, file: UploadFile = File(...)) -> Source:
    _require_notebook(notebook_id)
    payload = await file.read()
    if not payload:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "The uploaded file is empty.")
    filename = file.filename or "upload"
    try:
        return _ingestion.ingest_file(notebook_id, payload, filename)
    except IngestionError as exc:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, str(exc)) from exc
    except Exception as exc:  # never surface an opaque 500 to the user
        logging.getLogger("briefcase").exception("upload_failed: %s", filename)
        raise HTTPException(
            status.HTTP_422_UNPROCESSABLE_ENTITY,
            f"Could not process '{filename}': {type(exc).__name__}. "
            "See the server console for details.",
        ) from exc


@router.post(
    "/notebooks/{notebook_id}/sources/text",
    response_model=Source,
    status_code=status.HTTP_201_CREATED,
)
def add_text_source(notebook_id: str, payload: AddTextSource) -> Source:
    _require_notebook(notebook_id)
    try:
        return _ingestion.ingest_text(notebook_id, payload.title, payload.text)
    except IngestionError as exc:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, str(exc)) from exc


@router.post(
    "/notebooks/{notebook_id}/sources/url",
    response_model=Source,
    status_code=status.HTTP_201_CREATED,
)
def add_url_source(notebook_id: str, payload: AddUrlSource) -> Source:
    _require_notebook(notebook_id)
    try:
        return _ingestion.ingest_url(notebook_id, payload.url, payload.title)
    except IngestionError as exc:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, str(exc)) from exc


@router.get("/notebooks/{notebook_id}/sources/{source_id}/text")
def get_source_text(notebook_id: str, source_id: str) -> dict:
    result = repo.get_source_text(notebook_id, source_id)
    if result is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Source not found")
    title, text = result
    return {"source_id": source_id, "title": title, "text": text}


@router.delete(
    "/notebooks/{notebook_id}/sources/{source_id}", status_code=status.HTTP_204_NO_CONTENT
)
def delete_source(notebook_id: str, source_id: str) -> None:
    if not repo.delete_source(notebook_id, source_id):
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Source not found")


@router.post("/notebooks/{notebook_id}/sources/{source_id}/toggle")
def toggle_source(notebook_id: str, source_id: str, payload: dict) -> dict:
    enabled = bool((payload or {}).get("enabled", True))
    if not repo.toggle_source(notebook_id, source_id, enabled):
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Source not found")
    return {"source_id": source_id, "enabled": enabled}


@router.post("/notebooks/{notebook_id}/sources/{source_id}/refresh")
def refresh_source(notebook_id: str, source_id: str) -> dict:
    try:
        count = _ingestion.refresh_source(notebook_id, source_id)
    except IngestionError as exc:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, str(exc)) from exc
    return {"source_id": source_id, "chunk_count": count}


# ---------------- chat / studio ----------------


@router.post("/notebooks/{notebook_id}/chat", response_model=Answer)
def chat(notebook_id: str, payload: ChatRequest) -> Answer:
    _require_notebook(notebook_id)
    return _generator.answer(
        notebook_id,
        payload.question,
        source_ids=payload.source_ids or None,
        top_k=payload.top_k,
    )


@router.post("/notebooks/{notebook_id}/summary")
def summary(notebook_id: str, payload: dict | None = None) -> dict:
    _require_notebook(notebook_id)
    source_ids = (payload or {}).get("source_ids") or None
    return studio.summarize(notebook_id, source_ids)


@router.post("/notebooks/{notebook_id}/questions")
def questions(notebook_id: str, payload: dict | None = None) -> dict:
    _require_notebook(notebook_id)
    source_ids = (payload or {}).get("source_ids") or None
    return studio.suggested_questions(notebook_id, source_ids)


# ---------------- draft (sink) ----------------


@router.get("/notebooks/{notebook_id}/draft", response_model=Draft)
def get_draft(notebook_id: str) -> Draft:
    _require_notebook(notebook_id)
    return Draft(**repo.get_or_create_draft(notebook_id))


@router.put("/notebooks/{notebook_id}/draft", response_model=Draft)
def save_draft(notebook_id: str, payload: SaveDraft) -> Draft:
    _require_notebook(notebook_id)
    return Draft(**repo.save_draft(notebook_id, title=payload.title, body=payload.body))


@router.post("/notebooks/{notebook_id}/draft/command", response_model=DraftEdit)
def draft_command(notebook_id: str, payload: DraftCommand) -> DraftEdit:
    _require_notebook(notebook_id)
    body = repo.get_or_create_draft(notebook_id)["body"]
    source_ids = payload.source_ids or None
    if payload.mode in ("edit", "write"):
        edit = _author.edit(notebook_id, payload.command, body, source_ids=source_ids)
    elif payload.mode == "rewrite":
        edit = _author.rewrite(notebook_id, payload.command, payload.selection, body, source_ids=source_ids)
    elif payload.mode == "check":
        edit = _author.check(notebook_id, body, source_ids=source_ids)
    else:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, f"Unknown mode: {payload.mode}")

    # Persist body-changing edits with an undo snapshot; 'check' never changes the body.
    if payload.mode in ("edit", "write", "rewrite") and edit.body != body:
        repo.save_draft(notebook_id, body=edit.body, snapshot_note=f"before {payload.mode}")
    return edit


@router.post("/notebooks/{notebook_id}/draft/import", response_model=Draft)
async def import_into_draft(notebook_id: str, file: UploadFile = File(...)) -> Draft:
    """Seed the draft (sink) from an uploaded file: text, PDF, DOCX, PPTX, Markdown,
    HTML, code, or a scanned/image file (via OCR)."""
    _require_notebook(notebook_id)
    from briefcase.ingest.parsers import ParserError, parse

    payload = await file.read()
    if not payload:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "The uploaded file is empty.")
    try:
        doc = parse(payload, filename=file.filename or "import")
    except ParserError as exc:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, str(exc)) from exc
    except Exception as exc:
        raise HTTPException(
            status.HTTP_422_UNPROCESSABLE_ENTITY,
            f"Could not read '{file.filename}': {type(exc).__name__}.",
        ) from exc
    existing = repo.get_or_create_draft(notebook_id)["body"]
    new_body = f"{existing.rstrip()}\n\n{doc.text}".strip() if existing.strip() else doc.text
    return Draft(**repo.save_draft(notebook_id, body=new_body, snapshot_note="before import"))


def _safe_filename(name: str, ext: str) -> str:
    import re as _re

    stem = _re.sub(r'[^\w\- ]+', "", name).strip() or "draft"
    return f"{stem}.{ext}"


@router.get("/notebooks/{notebook_id}/draft/export.docx")
def export_draft_docx(notebook_id: str) -> Response:
    """Export the draft as a Word document (minimal Markdown -> docx conversion)."""
    import io
    import re

    from docx import Document

    _require_notebook(notebook_id)
    draft = repo.get_or_create_draft(notebook_id)

    def add_md(doc, text: str, style=None) -> None:
        para = doc.add_paragraph(style=style)
        for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
            run = para.add_run(part)
            if i % 2 == 1:
                run.bold = True

    document = Document()
    document.add_heading(draft["title"] or "Draft", level=0)
    for raw in draft["body"].split("\n"):
        line = raw.rstrip()
        if not line.strip():
            continue
        if line.startswith("### "):
            document.add_heading(line[4:], level=3)
        elif line.startswith("## "):
            document.add_heading(line[3:], level=2)
        elif line.startswith("# "):
            document.add_heading(line[2:], level=1)
        elif line.lstrip().startswith(("- ", "* ")):
            add_md(document, line.lstrip()[2:], style="List Bullet")
        else:
            add_md(document, line)

    buf = io.BytesIO()
    document.save(buf)
    filename = _safe_filename(draft["title"], "docx")
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/notebooks/{notebook_id}/draft/versions", response_model=list[DraftVersion])
def draft_versions(notebook_id: str) -> list[DraftVersion]:
    _require_notebook(notebook_id)
    return [DraftVersion(**v) for v in repo.list_draft_versions(notebook_id)]


@router.post("/notebooks/{notebook_id}/draft/restore", response_model=Draft)
def restore_draft(notebook_id: str, payload: dict) -> Draft:
    _require_notebook(notebook_id)
    version_id = (payload or {}).get("version_id")
    result = repo.restore_draft_version(notebook_id, version_id or "")
    if result is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Version not found")
    return Draft(**result)
